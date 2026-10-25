import os
import uuid
import base64
from datetime import datetime
from typing import Optional, List

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import nullslast

from database import get_db
import models
import ai_service
from auth import get_current_user
from routers.observations import get_allowed_project_ids

router = APIRouter(prefix="/api/meeting-minutes", tags=["meeting-minutes"])

# Roles allowed to create / edit / delete meeting minutes.
# "PSO and CST/admin" — interpreted as PSO plus Admin/SuperAdmin.
WRITE_ROLES = ["SuperAdmin", "Admin", "PSO"]

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)
MAX_FILE_BYTES = 25 * 1024 * 1024  # 25 MB
# Uploadable minutes formats. Legacy .doc is not supported (no reliable pure-Python parser).
ALLOWED_EXTS = {".pdf", ".docx", ".txt", ".csv", ".xlsx", ".xlsm", ".xls"}
# Scanned PDFs are summarised from page images — cap the pages sent to the model.
MAX_SCANNED_PAGES = 8


class MinutesCreate(BaseModel):
    title: str
    meeting_date: Optional[str] = None
    project_id: Optional[int] = None
    attendees: Optional[str] = None
    content: Optional[str] = None


class MinutesUpdate(BaseModel):
    title: Optional[str] = None
    meeting_date: Optional[str] = None
    project_id: Optional[int] = None
    attendees: Optional[str] = None
    content: Optional[str] = None


def _assert_can_write(user: models.User):
    if user.role not in WRITE_ROLES:
        raise HTTPException(403, "Only PSO and Admin users can manage meeting minutes")


def _assert_project_access(m: models.MeetingMinutes, user: models.User):
    allowed = get_allowed_project_ids(user)
    if allowed is not None and m.project_id is not None and m.project_id not in allowed:
        raise HTTPException(403, "Access denied to these meeting minutes")


def mom_to_dict(m: models.MeetingMinutes) -> dict:
    return {
        "id": m.id,
        "title": m.title,
        "meeting_date": m.meeting_date,
        "project_id": m.project_id,
        "project_name": m.project.name if m.project else None,
        "attendees": m.attendees,
        "entry_type": m.entry_type,
        "content": m.content,
        "file_path": m.file_path,
        "file_name": m.file_name,
        "ai_summary": m.ai_summary,
        "ai_summary_at": m.ai_summary_at.isoformat() if m.ai_summary_at else None,
        "ai_summary_by": m.ai_summary_by,
        "ai_summary_by_name": m.ai_summary_user.name if m.ai_summary_user else None,
        "created_by": m.created_by,
        "created_by_name": m.creator.name if m.creator else None,
        "created_at": m.created_at.isoformat() if m.created_at else None,
        "updated_at": m.updated_at.isoformat() if m.updated_at else None,
    }


def _extract_pdf_text(abs_path: str) -> str:
    """Extract text from a PDF for summarisation. Returns '' if nothing could be read."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise HTTPException(500, "PDF support is not installed on the server (pypdf missing).")
    try:
        reader = PdfReader(abs_path)
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _pdf_page_images(abs_path: str) -> list:
    """Render PDF pages to JPEG for scanned (image-only) PDFs. Returns [(media_type, base64)]."""
    try:
        import pymupdf
    except ImportError:
        raise HTTPException(500, "Scanned-PDF support is not installed on the server (pymupdf missing).")
    out = []
    try:
        with pymupdf.open(abs_path) as doc:
            for page in list(doc)[:MAX_SCANNED_PAGES]:
                # ~1600px on the long edge is plenty for the model to read the page.
                zoom = 1600 / max(page.rect.width, page.rect.height)
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
                data = pix.tobytes("jpeg", jpg_quality=80)
                out.append(("image/jpeg", base64.standard_b64encode(data).decode("utf-8")))
    except Exception:
        return []
    return out


def _extract_docx_text(abs_path: str) -> str:
    import docx
    doc = docx.Document(abs_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Minutes are often laid out in tables (action-item grids etc.).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_excel_text(abs_path: str, ext: str) -> str:
    parts = []
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(abs_path)
        for sheet in book.sheets():
            parts.append(f"# Sheet: {sheet.name}")
            for r in range(sheet.nrows):
                cells = [str(v).strip() for v in sheet.row_values(r)]
                if any(cells):
                    parts.append(" | ".join(cells))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(abs_path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            parts.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    parts.append(" | ".join(cells))
        wb.close()
    return "\n".join(parts).strip()


def _extract_text_file(abs_path: str) -> str:
    with open(abs_path, "rb") as fp:
        raw = fp.read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return ""


def _extract_file_text(abs_path: str) -> str:
    """Extract text from an uploaded minutes file. Returns '' if nothing could be read."""
    ext = os.path.splitext(abs_path)[1].lower()
    if ext == ".pdf":
        return _extract_pdf_text(abs_path)
    try:
        if ext == ".docx":
            return _extract_docx_text(abs_path)
        if ext in (".xlsx", ".xlsm", ".xls"):
            return _extract_excel_text(abs_path, ext)
        if ext in (".txt", ".csv"):
            return _extract_text_file(abs_path)
    except Exception:
        return ""
    return ""


def _scoped_query(db: Session, user: models.User):
    q = db.query(models.MeetingMinutes)
    allowed = get_allowed_project_ids(user)
    if allowed is not None:
        q = q.filter(
            (models.MeetingMinutes.project_id.in_(allowed))
            | (models.MeetingMinutes.project_id.is_(None))
        )
    return q


@router.get("/")
def list_minutes(
    project_id: List[int] = Query(default=[]),
    search: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user: models.User = Depends(get_current_user),
):
    q = _scoped_query(db, user)
    if project_id:
        q = q.filter(models.MeetingMinutes.project_id.in_(project_id))
    if search:
        like = f"%{search}%"
        q = q.filter(
            (models.MeetingMinutes.title.ilike(like))
            | (models.MeetingMinutes.attendees.ilike(like))
            | (models.MeetingMinutes.content.ilike(like))
        )
    rows = q.order_by(
        nullslast(models.MeetingMinutes.meeting_date.desc()),
        models.MeetingMinutes.created_at.desc(),
    ).all()
    return {"minutes": [mom_to_dict(m) for m in rows]}


@router.get("/{mom_id}")
def get_minutes(mom_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    m = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == mom_id).first()
    if not m:
        raise HTTPException(404, "Meeting minutes not found")
    _assert_project_access(m, user)
    return mom_to_dict(m)


@router.post("/", status_code=201)
def create_text_minutes(body: MinutesCreate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    _assert_can_write(user)
    if not body.title or not body.title.strip():
        raise HTTPException(400, "Title is required")
    if not body.content or not body.content.strip():
        raise HTTPException(400, "Minutes content is required")
    m = models.MeetingMinutes(
        title=body.title.strip(),
        meeting_date=body.meeting_date,
        project_id=body.project_id,
        attendees=body.attendees,
        entry_type="text",
        content=body.content.strip(),
        created_by=user.id,
    )
    db.add(m)
    db.commit()
    db.refresh(m)
    return mom_to_dict(m)


@router.post("/upload", status_code=201)
async def upload_file_minutes(
    title: str = Form(...),
    meeting_date: Optional[str] = Form(None),
    project_id: Optional[int] = Form(None),
    attendees: Optional[str] = Form(None),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: models.User = Depends(get_current_user),
):
    _assert_can_write(user)
    if not title or not title.strip():
        raise HTTPException(400, "Title is required")
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext == ".doc":
        raise HTTPException(400, "Legacy .doc files are not supported — please save it as .docx or PDF")
    if ext not in ALLOWED_EXTS:
        raise HTTPException(400, "Only PDF, Word (.docx), Excel (.xlsx/.xls), CSV or text files are allowed")
    content = await file.read()
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(400, "File is too large (max 25 MB)")

    fname = f"{uuid.uuid4().hex}{ext}"
    with open(os.path.join(UPLOAD_DIR, fname), "wb") as fp:
        fp.write(content)

    m = models.MeetingMinutes(
        title=title.strip(),
        meeting_date=meeting_date,
        project_id=project_id,
        attendees=attendees,
        entry_type="pdf" if ext == ".pdf" else "file",
        file_path=fname,
        file_name=file.filename or fname,
        created_by=user.id,
    )
    db.add(m)
    db.commit()
    db.refresh(m)
    return mom_to_dict(m)


@router.put("/{mom_id}")
def update_minutes(mom_id: int, body: MinutesUpdate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    m = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == mom_id).first()
    if not m:
        raise HTTPException(404, "Meeting minutes not found")
    _assert_can_write(user)
    _assert_project_access(m, user)

    data = body.model_dump(exclude_unset=True)
    if "title" in data:
        if not data["title"] or not data["title"].strip():
            raise HTTPException(400, "Title cannot be empty")
        data["title"] = data["title"].strip()
    for field, val in data.items():
        setattr(m, field, val)
    m.updated_at = datetime.now()
    db.commit()
    db.refresh(m)
    return mom_to_dict(m)


@router.post("/{mom_id}/ai-summary")
def generate_ai_summary(mom_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    m = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == mom_id).first()
    if not m:
        raise HTTPException(404, "Meeting minutes not found")
    _assert_project_access(m, user)
    if not ai_service.is_enabled(db):
        raise HTTPException(503, "AI summaries are turned off. Enable them in Admin → AI Summary.")
    if not ai_service.is_configured(db):
        raise HTTPException(503, "AI summary is not configured. Add an API key in Admin → AI Summary.")

    # Gather the source: text extracted from the uploaded file, or the typed content.
    # Scanned PDFs have no text layer, so their pages are sent to the model as images.
    images = []
    if m.entry_type in ("pdf", "file") and m.file_path:
        abs_path = os.path.join(UPLOAD_DIR, m.file_path)
        if not os.path.exists(abs_path):
            raise HTTPException(404, "The uploaded file is missing on the server.")
        source = _extract_file_text(abs_path)
        if len(source) < 20 and abs_path.lower().endswith(".pdf"):
            images = _pdf_page_images(abs_path)
        if not source and not images:
            raise HTTPException(422, "Could not read any content from this file.")
    else:
        source = m.content or ""
        if not source.strip():
            raise HTTPException(422, "There is no text to summarise.")

    try:
        summary = ai_service.summarize_minutes(source, db, images=images)
    except Exception:
        raise HTTPException(502, "Failed to generate the AI summary. Please try again.")

    m.ai_summary = summary
    m.ai_summary_at = datetime.now()
    m.ai_summary_by = user.id
    db.commit()
    db.refresh(m)
    return {
        "ai_summary": m.ai_summary,
        "ai_summary_at": m.ai_summary_at.isoformat() if m.ai_summary_at else None,
        "ai_summary_by": m.ai_summary_by,
        "ai_summary_by_name": user.name,
    }


@router.delete("/{mom_id}")
def delete_minutes(mom_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    m = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == mom_id).first()
    if not m:
        raise HTTPException(404, "Meeting minutes not found")
    is_super_admin = user.role == "SuperAdmin"
    is_owner_writer = m.created_by == user.id and user.role in WRITE_ROLES
    if not (is_super_admin or is_owner_writer or user.role == "Admin"):
        raise HTTPException(403, "Only an Admin, SuperAdmin or the creator can delete these minutes")
    if m.file_path:
        try:
            os.remove(os.path.join(UPLOAD_DIR, m.file_path))
        except OSError:
            pass
    db.delete(m)
    db.commit()
    return {"success": True}
